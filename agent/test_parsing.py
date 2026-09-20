from app.parsing import parse_to_markdown


def test_text_parser_returns_markdown():
    markdown, parser = parse_to_markdown("guide.md", b"# VPN\n\nReset the client cache.")
    assert parser == "text"
    assert markdown.startswith("# VPN")


def test_unknown_extension_is_rejected():
    try:
        parse_to_markdown("archive.zip", b"bytes")
    except ValueError as exc:
        assert "不支持" in str(exc)
    else:
        raise AssertionError("unsupported file should be rejected")


def test_docx_keeps_paragraphs_and_tables(tmp_path):
    """docx 解析按文档顺序保留段落与表格（旧实现只读 paragraphs，表格内容全部丢失）。"""
    import pytest
    docx = pytest.importorskip("docx")
    from docx import Document

    document = Document()
    document.add_paragraph("VPN 配置指南")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "server"
    table.cell(1, 1).text = "vpn.example.com"
    path = tmp_path / "guide.docx"
    document.save(path)

    markdown, parser = parse_to_markdown("guide.docx", path.read_bytes())
    assert parser == "docx"
    assert "VPN 配置指南" in markdown
    assert "| 参数 | 值 |" in markdown
    assert "| server | vpn.example.com |" in markdown
